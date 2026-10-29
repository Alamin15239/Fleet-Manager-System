from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

OUT = Path("pitch-deck/Fleet_Manager_Pro_Startup_Pitch_Deck.pptx")
OUT.parent.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(11, 22, 42)
PANEL = RGBColor(20, 37, 64)
BLUE = RGBColor(54, 131, 255)
CYAN = RGBColor(53, 211, 205)
WHITE = RGBColor(246, 249, 255)
MUTED = RGBColor(171, 187, 209)
ORANGE = RGBColor(255, 176, 91)
GREEN = RGBColor(104, 224, 173)

def textbox(slide, x, y, w, h, text, size=20, color=WHITE, bold=False, font="Aptos", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02); tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    if align is not None: p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box

def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None: shape.line.color.rgb = line
    if radius:
        try: shape.adjustments[0] = 0.12
        except Exception: pass
    return shape

def base(title, eyebrow="FLEET MANAGER PRO", num=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.12, 7.5, BLUE)
    textbox(slide, 0.65, 0.35, 8, 0.25, eyebrow.upper(), 10, CYAN, True)
    textbox(slide, 0.65, 0.82, 12, 0.7, title, 30, WHITE, True)
    rect(slide, 0.65, 1.62, 1.0, 0.055, BLUE)
    textbox(slide, 0.65, 7.12, 9, 0.18, "FLEET MAINTENANCE • TIRES • WORKSHOP • DOCUMENTS", 8, MUTED)
    if num is not None: textbox(slide, 12.15, 7.08, 0.5, 0.25, f"{num:02d}", 10, MUTED, True, align=PP_ALIGN.RIGHT)
    return slide

def bullets(slide, items, x=0.9, y=2.05, w=7.0, size=20, gap=0.72):
    for i, item in enumerate(items):
        yy = y + i*gap
        rect(slide, x, yy+0.12, 0.13, 0.13, CYAN, True)
        textbox(slide, x+0.28, yy, w-0.3, gap-0.08, item, size, WHITE)

def card(slide, x, y, w, h, label, desc, accent=BLUE):
    rect(slide, x, y, w, h, PANEL, True)
    rect(slide, x, y, 0.07, h, accent, True)
    textbox(slide, x+0.25, y+0.2, w-0.45, 0.38, label, 17, WHITE, True)
    textbox(slide, x+0.25, y+0.68, w-0.45, h-0.8, desc, 12, MUTED)

# 1 cover
s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
rect(s, 8.7, 0, 4.63, 7.5, PANEL)
for i, (x,y,w,h,c) in enumerate([(9.25,1.1,3.2,0.85,BLUE),(9.75,2.25,2.7,0.85,CYAN),(9.25,3.4,3.2,0.85,RGBColor(39,78,130)),(9.75,4.55,2.7,0.85,RGBColor(39,110,122))]):
    rect(s,x,y,w,h,c,True)
    textbox(s,x+0.18,y+0.2,w-0.35,0.35,["VEHICLES","MAINTENANCE","TIRE RECORDS","DOCUMENTS"][i],12,WHITE,True)
textbox(s,0.8,0.72,6.8,0.3,"STARTUP PITCH • OCTOBER 2026",11,CYAN,True)
textbox(s,0.8,1.55,7.5,1.5,"Fleet Manager\nPro",42,WHITE,True)
textbox(s,0.85,3.45,6.7,0.85,"One workspace for fleet maintenance and operational records.",22,MUTED)
rect(s,0.85,4.7,1.25,0.06,BLUE)
textbox(s,0.85,5.05,6.5,0.35,"Founder: Alamin Khan",15,WHITE,True)
textbox(s,0.85,5.52,6.5,0.35,"fleetinspections.tech",14,CYAN)
textbox(s,0.85,6.85,7.0,0.25,"EARLY-STAGE • PRE-REVENUE • CUSTOMER VALIDATION",9,MUTED,True)

# 2 thesis
s=base("The operating problem we want to solve",num=2)
textbox(s,0.85,2.0,5.7,0.7,"Operational records can be fragmented.",25,WHITE,True)
bullets(s,["Vehicle history in one place, tire records in another","Workshop jobs and breakdowns tracked separately","Document expiry dates can be easy to miss"],0.9,3.0,6.1,17,0.8)
card(s,7.5,2.15,4.8,1.25,"THE HYPOTHESIS","A shared operational view may reduce follow-up friction.",CYAN)
textbox(s,7.55,3.8,4.6,1.3,"These are customer-discovery hypotheses—not measured market findings.",15,ORANGE,True)

# 3 solution
s=base("A centralized workspace for fleet operations",num=3)
card(s,.8,2.05,3.75,1.55,"Fleet & vehicles","Vehicle records and maintenance history.",BLUE)
card(s,4.8,2.05,3.75,1.55,"Tires & movement","Inventory, serials, transfers and remarks.",CYAN)
card(s,8.8,2.05,3.75,1.55,"Workshop & breakdowns","Job cards and incident records.",ORANGE)
card(s,2.8,4.15,3.75,1.55,"Documents","Issue/expiry dates and configurable alerts.",GREEN)
card(s,6.8,4.15,3.75,1.55,"Reports & integrations","Operational reporting; validate sync reliability live.",BLUE)

# 4 workflow
s=base("From record capture to follow-up",num=4)
steps=[("01","RECORD","Create or update a vehicle record"),("02","LOG","Add a service, tire, job or breakdown event"),("03","REVIEW","See relevant history in its module"),("04","FOLLOW UP","Use recorded dates to plan next action")]
for i,(n,t,d) in enumerate(steps):
    x=.8+i*3.12
    rect(s,x,2.45,2.65,2.5,PANEL,True)
    textbox(s,x+.22,2.7,1.0,.4,n,20,CYAN,True)
    textbox(s,x+.22,3.35,2.2,.4,t,15,WHITE,True)
    textbox(s,x+.22,3.9,2.15,.85,d,13,MUTED)
    if i<3: textbox(s,x+2.72,3.35,.35,.4,"→",20,BLUE,True)

# 5 modules
s=base("Designed around practical fleet workflows",num=5)
modules=[("Preventive maintenance","Service history and maintenance planning"),("Breakdown management","Incident records and review"),("Tire inventory","Serials, transfers and damage notes"),("Workshop job cards","Structured workshop work records"),("Expiry alerts","Configurable vehicle-document reminders"),("Excel / OneDrive","Integration workflow; reliability to be validated")]
for i,(a,b) in enumerate(modules):
    x=.8+(i%3)*4.15; y=2.0+(i//3)*2.15
    card(s,x,y,3.7,1.65,a,b,[BLUE,CYAN,ORANGE,GREEN,BLUE,CYAN][i])

# 6 users
s=base("Who we plan to serve first",num=6)
card(s,.9,2.25,3.65,2.25,"TRUCKING FLEETS","Teams coordinating vehicles, service events and records.",BLUE)
card(s,4.85,2.25,3.65,2.25,"LOGISTICS COMPANIES","Transport operators needing a consistent operational view.",CYAN)
card(s,8.8,2.25,3.65,2.25,"INDEPENDENT WORKSHOPS","Teams tracking workshop jobs and vehicle visits.",ORANGE)
textbox(s,1.0,5.15,11.2,.5,"Initial geography: Saudi Arabia • International expansion considered after validation",16,MUTED,True,align=PP_ALIGN.CENTER)
textbox(s,1.0,5.8,11.2,.35,"These are target segments, not confirmed customers.",12,ORANGE,False,align=PP_ALIGN.CENTER)

# 7 product
s=base("Product direction: operational clarity, not more spreadsheets",num=7)
textbox(s,.95,2.05,5.1,.7,"Fleet Manager Pro",27,WHITE,True)
textbox(s,.95,2.85,5.2,1.1,"A web-based product focused on bringing maintenance, tire, breakdown, workshop and document records into one workspace.",18,MUTED)
card(s,7.0,2.05,5.0,1.1,"LIVE PRODUCT","fleetinspections.tech",CYAN)
card(s,7.0,3.45,5.0,1.1,"NEXT PROOF POINT","Run guided pilots against real fleet workflows.",BLUE)
card(s,7.0,4.85,5.0,1.1,"QUALITY BAR","Validate imports, sync, date extraction and alerts.",ORANGE)

# 8 traction
s=base("Current traction — an honest baseline",num=8)
metrics=[("2","Unique users","Last 30 days"),("$0","Revenue","Last 30 days"),("No","Registered company","As supplied by founder"),("0","Co-founders","As supplied by founder")]
for i,(v,l,d) in enumerate(metrics):
    x=.8+i*3.15
    rect(s,x,2.35,2.75,2.15,PANEL,True)
    textbox(s,x+.2,2.7,2.3,.65,v,32,[CYAN,BLUE,ORANGE,GREEN][i],True,align=PP_ALIGN.CENTER)
    textbox(s,x+.2,3.45,2.3,.35,l,15,WHITE,True,align=PP_ALIGN.CENTER)
    textbox(s,x+.2,3.92,2.3,.35,d,10,MUTED,align=PP_ALIGN.CENTER)
textbox(s,1.0,5.35,11.2,.5,"Stage: early-stage / pre-revenue. No paying customers or verified savings are claimed.",14,ORANGE,True,align=PP_ALIGN.CENTER)

# 9 business model
s=base("Business model hypothesis",num=9)
textbox(s,.95,2.0,5.5,.7,"Subscription tiers by fleet size and workflow needs.",23,WHITE,True)
card(s,.95,3.05,3.55,1.8,"STARTER","Core records for a small fleet.",BLUE)
card(s,4.9,3.05,3.55,1.8,"GROWTH","Broader workflows and reporting.",CYAN)
card(s,8.85,3.05,3.55,1.8,"BUSINESS","Larger-team requirements and controls.",ORANGE)
textbox(s,1.0,5.5,11.0,.4,"Pricing and willingness to pay are not yet validated.",15,ORANGE,True,align=PP_ALIGN.CENTER)

# 10 go to market
s=base("Go-to-market: learn with guided pilots",num=10)
bullets(s,["Interview fleet operators about current workflows","Recruit a small number of pilot users","Agree success criteria before each pilot","Measure usage, task completion and record completeness","Ask for willingness-to-pay feedback before pricing decisions"],.95,2.05,10.9,18,.83)
textbox(s,1.0,6.2,10.8,.3,"Goal: evidence of repeatable value before scaling acquisition.",13,CYAN,True)

# 11 roadmap
s=base("Roadmap: validate, harden, then expand",num=11)
road=[("NOW","Core workflow validation","Test maintenance, tires, breakdowns, jobs and documents."),("NEXT","Reliability & onboarding","Improve error handling, imports, date extraction and integration behavior."),("THEN","Pilot-to-paid readiness","Set support, data handling, backups and success criteria."),("LATER","Demand-led expansion","Expand where recurring customer demand supports it.")]
for i,(t,a,b) in enumerate(road):
    y=2.0+i*1.05
    rect(s,.95,y,1.3,.55,[BLUE,CYAN,ORANGE,GREEN][i],True)
    textbox(s,1.03,y+.12,1.1,.25,t,11,NAVY,True,align=PP_ALIGN.CENTER)
    textbox(s,2.55,y-.02,3.0,.35,a,16,WHITE,True)
    textbox(s,5.6,y-.02,6.5,.65,b,13,MUTED)

# 12 risks
s=base("Risks we are actively validating",num=12)
card(s,.9,2.1,5.55,1.55,"PRODUCT RELIABILITY","Test date extraction, imports, sync and alert behavior.",ORANGE)
card(s,6.85,2.1,5.55,1.55,"CUSTOMER DEMAND","Interview users and test willingness to pay.",BLUE)
card(s,.9,4.15,5.55,1.55,"TRUST & OPERATIONS","Clarify data handling, support and backup practices.",CYAN)
card(s,6.85,4.15,5.55,1.55,"LEGAL READINESS","Confirm lawful business operations before paid scaling.",GREEN)

# 13 differentiation
s=base("Our focus: connected records for day-to-day operations",num=13)
textbox(s,.95,2.1,5.2,.55,"The product thesis",24,WHITE,True)
bullets(s,["Connect adjacent workflows in one workspace","Keep operational history close to the relevant record","Make follow-up easier through recorded dates and alerts"],.95,3.0,6.1,16,.85)
rect(s,7.55,2.15,4.55,3.35,PANEL,True)
for i,(t,c) in enumerate([("FLEET",BLUE),("TIRES",CYAN),("WORKSHOP",ORANGE),("DOCUMENTS",GREEN)]):
    rect(s,7.95,2.55+i*.65,3.75,.42,c,True)
    textbox(s,8.12,2.63+i*.65,3.4,.2,t,11,WHITE,True)
textbox(s,7.65,5.75,4.4,.45,"Positioning hypothesis—not a validated moat.",11,ORANGE,True)

# 14 ask
s=base("What we are looking for now",num=14)
asks=[("PILOT USERS","Fleet operators willing to test real workflows"),("INTRODUCTIONS","Connections to transport and workshop teams"),("MENTORSHIP","Product, reliability and go-to-market feedback")]
for i,(a,b) in enumerate(asks):
    card(s,.95+i*4.15,2.55,3.65,2.25,a,b,[BLUE,CYAN,ORANGE][i])
textbox(s,1.0,5.65,11.0,.5,"Success means learning what customers repeatedly use—and what they will pay for.",14,MUTED,True,align=PP_ALIGN.CENTER)

# 15 close
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=NAVY
rect(s,0,0,.12,7.5,BLUE)
textbox(s,.9,.75,7,.3,"FLEET MANAGER PRO",11,CYAN,True)
textbox(s,.9,1.65,8.5,1.2,"Let's make fleet\noperations clearer.",38,WHITE,True)
textbox(s,.95,4.3,7.2,.45,"Founder: Alamin Khan",17,MUTED,True)
textbox(s,.95,4.9,7.2,.45,"https://fleetinspections.tech",18,CYAN,True)
rect(s,9.1,1.65,3.0,3.0,PANEL,True)
textbox(s,9.45,2.2,2.3,.55,"FLEET",17,WHITE,True,align=PP_ALIGN.CENTER)
textbox(s,9.45,2.9,2.3,.55,"MANAGER",17,CYAN,True,align=PP_ALIGN.CENTER)
textbox(s,9.45,3.6,2.3,.55,"PRO",25,WHITE,True,align=PP_ALIGN.CENTER)
textbox(s,.95,6.85,10.5,.25,"Early-stage • Pre-revenue • Prepared October 2026",9,MUTED)

prs.save(OUT)
print(f"Created {OUT} ({len(prs.slides)} slides)")
